"""Node 서버에서 호출하는 엑셀 처리 서비스.

stdin 으로 JSON 요청을 받고, stdout 으로 JSON 응답을 돌려준다.
    요청: {"action": "read" | "write", ...}
    응답: {"ok": true, "data": ...} 또는 {"ok": false, "error": "..."}
"""
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"


def resolve_path(file_name):
    # data/ 폴더 밖으로 나가는 경로는 막는다
    path = (DATA_DIR / file_name).resolve()
    if DATA_DIR.resolve() not in path.parents:
        raise ValueError(f"허용되지 않은 경로입니다: {file_name}")
    return path


def read_excel(file, sheet=None):
    workbook = openpyxl.load_workbook(resolve_path(file), data_only=True)
    ws = workbook[sheet] if sheet else workbook.active
    rows = [list(row) for row in ws.iter_rows(values_only=True)]
    return {"sheet": ws.title, "rows": rows}


def write_excel(file, rows, sheet=None):
    path = resolve_path(file)
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = openpyxl.Workbook()
    ws = workbook.active
    if sheet:
        ws.title = sheet
    for row in rows:
        ws.append(row)

    workbook.save(path)
    return {"file": str(path.relative_to(ROOT)), "rowCount": len(rows)}


ACTIONS = {
    "read": read_excel,
    "write": write_excel,
}


def main():
    try:
        # utf-8-sig: BOM이 붙은 입력(PowerShell 파이프 등)도 허용
        request = json.loads(sys.stdin.buffer.read().decode("utf-8-sig"))
        action = request.pop("action", None)
        if action not in ACTIONS:
            raise ValueError(f"알 수 없는 action: {action}")
        result = {"ok": True, "data": ACTIONS[action](**request)}
    except Exception as e:
        result = {"ok": False, "error": str(e)}

    print(json.dumps(result, ensure_ascii=False, default=str))


if __name__ == "__main__":
    main()
